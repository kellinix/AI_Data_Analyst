"""Compare public Kaggle CSVs with Zephyr's real cleaning and analysis path.

Download the CC0 ihelon/coffee-sales and Apache-2.0
pimmadachirawat/supermarket-sales datasets separately. Source files stay out
of Git. Run from backend with explicit paths, for example::

    python -m evals.reliability.run_kaggle_eval --coffee ../tmp/kaggle/coffee/index_1.csv \
        --supermarket "../tmp/kaggle/supermarket/Supermarket Sales Data.csv"

AI is disabled so this gate tests measured results, not generated prose.
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import csv
import json
import logging
import os
import re
import sys
import zlib
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from evals.calibration.run_eval import _bootstrap_env, _OfflineAIService

CASES = {
    "coffee": {"measure": "money", "group": "cash_type"},
    "supermarket": {"measure": "Total", "group": "Branch"},
}


def source_facts(path: Path, measure: str, group: str) -> dict:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    totals: dict[str, Decimal] = {}
    counts: Counter[str] = Counter()
    missing = 0
    total = Decimal(0)
    for row in rows:
        value = (row.get(measure) or "").strip()
        key = (row.get(group) or "").strip()
        counts[key] += 1
        if not value:
            missing += 1
            continue
        amount = Decimal(value.replace(",", ""))
        total += amount
        totals[key] = totals.get(key, Decimal(0)) + amount
    return {
        "rows": len(rows), "measure": measure, "group": group,
        "total": str(total), "missing_measure": missing,
        "group_counts": dict(counts),
        "group_totals": {key: str(value) for key, value in totals.items()},
    }


async def run_case(name: str, path: Path, workdir: Path, *, with_ai: bool) -> dict:
    print(f"{name}: loading source", file=sys.stderr, flush=True)
    import polars as pl

    from app.analytics.histogram_presentation import present_histogram
    from app.analytics.labels import derive_label
    from app.analytics.record_context import record_context
    from app.analytics.report_relevance import useful_charts
    from app.api.v1.endpoints.analyses import _cleaning_options
    from app.services.analysis_engine import compute_analysis
    from app.services.export_charts import normalise_charts
    from app.services.export_service import build_excel_report, build_pdf_report
    from app.services.file_processor import FileProcessor
    from app.services.report_language import present_insight, present_metadata

    config = CASES[name]
    source = source_facts(path, **config)
    target = workdir / f"{name}.cleaned.parquet"
    processor = FileProcessor()
    print(f"{name}: cleaning {source['rows']} rows", file=sys.stderr, flush=True)
    profile = await processor.clean_file(
        str(path), ".csv", str(target), str(target.with_suffix(".csv")),
        _cleaning_options({"mode": "clean"}),
    )
    frame = pl.read_parquet(target)
    print(f"{name}: analysing {frame.height} rows", file=sys.stderr, flush=True)
    if with_ai:
        from app.services.ai_service import AIService

        ai_service = AIService()
    else:
        ai_service = _OfflineAIService()
    computed = await compute_analysis(
        file_processor=processor, ai_service=ai_service,
        storage_path=str(target), file_name=path.name, analysis_id=f"kaggle-{name}",
        upload_context={"cleaning": {"enabled": True, "report": profile["cleaning_report"]}},
    )
    print(f"{name}: comparing results", file=sys.stderr, flush=True)
    measure_key = config["measure"].lower().replace(" ", "_")
    stats = computed["statistics"]["numeric_stats"].get(measure_key)
    actual_total = Decimal(str(stats["total"])) if stats and stats.get("total") is not None else None
    expected_total = Decimal(source["total"])
    failures = []
    if computed["row_count"] != source["rows"] or frame.height != source["rows"]:
        failures.append("Source rows changed during cleaning or analysis")
    if actual_total is None or abs(actual_total - expected_total) > Decimal("0.01"):
        failures.append(f"{config['measure']} total differs from source")
    if measure_key in frame.columns and frame[measure_key].null_count() != source["missing_measure"]:
        failures.append(f"{config['measure']} blanks changed during cleaning")
    group_key = config["group"].lower().replace(" ", "_")
    group_chart = next(
        (chart for chart in computed["charts"] if chart.get("type") == "bar"
         and chart.get("xAxis") == measure_key and chart.get("yAxis") == group_key),
        None,
    )
    displayed_totals: dict[str, Decimal] = {}
    if group_chart is None:
        failures.append(f"No {config['measure']} by {config['group']} chart")
    else:
        option = group_chart.get("echarts_option") or {}
        labels = (option.get("yAxis") or {}).get("data") or []
        values = ((option.get("series") or [{}])[0]).get("data") or []
        displayed_totals = {str(label): Decimal(str(value)) for label, value in zip(labels, values, strict=True)}
        expected_groups = {key: Decimal(value) for key, value in source["group_totals"].items()}
        if displayed_totals.keys() != expected_groups.keys() or any(
            abs(displayed_totals.get(key, Decimal(0)) - expected) > Decimal("0.01")
            for key, expected in expected_groups.items()
        ):
            failures.append(f"{config['measure']} by {config['group']} chart differs from source")
    headline = next((kpi for kpi in computed["kpis"] if kpi.get("column") == measure_key), None)
    if headline is None or abs(Decimal(str(headline["value"])) - expected_total) > Decimal("0.01"):
        failures.append(f"{config['measure']} headline differs from source")
    context = record_context(computed["statistics"])
    if not context or context.get("kind") != "sales_transaction":
        failures.append("Sales rows are not given a plain-language record name")
    if name == "supermarket" and any(chart.get("type") == "scatter" for chart in computed["charts"]):
        failures.append("Invoice arithmetic was presented as a relationship insight")
    # Use the same chart presentation and export functions as the HTTP routes.
    # The local gate tests the files themselves, not just the in-memory charts.
    export_charts = [present_histogram(chart, context) for chart in computed["charts"]]
    visible_charts = normalise_charts(useful_charts(export_charts))
    metadata = {**computed["statistics"], "forecasts": computed["forecasts"],
                "anomalies": computed["anomalies"]}
    _, summary = present_metadata(metadata, computed["ai_result"].get("executive_summary") or "")
    labels = {item["name"]: item["label"] for item in computed["semantic_display"]["columns"]}
    export_insights = []
    for kpi in computed["kpis"]:
        label = labels.get(kpi["column"], derive_label(kpi["column"]))
        title = label if kpi.get("is_total", True) else f"Avg {label}"
        if kpi.get("is_percent"):
            title = f"{label} Rate"
        export_insights.append({
            "type": "summary", "title": title,
            "data": {"value": kpi["value"], "is_currency": kpi.get("is_currency", False),
                     "currency": kpi.get("currency"), "currency_symbol": kpi.get("currency_symbol"),
                     "column": kpi["column"], "is_total": kpi.get("is_total", True)},
        })
    export_insights.extend(present_insight(item, metadata) for item in (computed["ai_result"].get("insights") or []))
    for item in computed["recommendations"]:
        stored = {**item, "type": "recommendation", "data": {
            "problem": item.get("problem"), "evidence": item.get("evidence"),
            "expected_impact": item.get("expected_impact"),
            "confidence_source": item.get("confidence_source"),
            "confidence_method": item.get("confidence_method"),
            **(item.get("data") or {}),
        }}
        export_insights.append(present_insight(stored, metadata))
    export_args = {
        "name": f"Kaggle {name}", "row_count": computed["row_count"],
        "column_count": len(computed["statistics"]["schema"]),
        "summary": summary, "insights": export_insights, "charts": export_charts,
    }
    pdf_path = workdir / f"{name}.pdf"
    xlsx_path = workdir / f"{name}.xlsx"
    print(f"{name}: exporting PDF and Excel", file=sys.stderr, flush=True)
    pdf_bytes = build_pdf_report(**export_args)
    plain_pdf_bytes = build_pdf_report(**{**export_args, "charts": []})
    xlsx_bytes = build_excel_report(**export_args)
    pdf_path.write_bytes(pdf_bytes)
    xlsx_path.write_bytes(xlsx_bytes)
    if not pdf_bytes.startswith(b"%PDF-") or len(pdf_bytes) < 5_000:
        failures.append("PDF export is missing or too small")
    if visible_charts and len(pdf_bytes) < len(plain_pdf_bytes) + 2_500:
        failures.append("PDF charts do not add expected visual drawing data")
    # ReportLab emits ASCII85 + zlib page streams; inspect their chart labels
    # and drawing commands without requiring another PDF parser at evaluation time.
    page_streams = [
        zlib.decompress(base64.a85decode(stream.strip(), adobe=True))
        for stream in re.findall(rb"stream\n(.*?)endstream", pdf_bytes, re.S)
    ]
    pdf_content = b"\n".join(page_streams)
    missing_pdf_titles = [chart.title for chart in visible_charts
                          if chart.title.encode("latin-1", "replace") not in pdf_content]
    if missing_pdf_titles:
        failures.append(f"PDF is missing chart titles: {missing_pdf_titles}")
    if visible_charts and b"Visual Charts" not in pdf_content:
        failures.append("PDF is missing the visual chart section")
    if b"First recorded point" in pdf_content:
        failures.append("PDF still shows the unwanted first-point label")
    first_month = "Mar 2024" if name == "coffee" else "Jan 2019"
    if first_month.encode() not in pdf_content:
        failures.append("PDF time chart does not show month and year")
    workbook = load_workbook(xlsx_path, data_only=False)
    exported_chart_count = len(workbook["Charts"]._charts) if "Charts" in workbook else 0
    if exported_chart_count != len(visible_charts):
        failures.append(f"Excel has {exported_chart_count} charts; expected {len(visible_charts)}")
    headline_export_title = next((item["title"] for item in export_insights
                                  if item.get("type") == "summary" and
                                  (item.get("data") or {}).get("column") == measure_key), None)
    if "Key Metrics" not in workbook or not any(
        row[0].value == headline_export_title
        and isinstance(row[1].value, int | float)
        and abs(Decimal(str(row[1].value)) - expected_total) <= Decimal("0.01")
        for row in workbook["Key Metrics"].iter_rows(min_row=2, max_col=2)
    ):
        failures.append("Excel headline KPI differs from source")
    if group_chart is not None and "Chart Data" in workbook:
        group_title = group_chart.get("title")
        sheet = workbook["Chart Data"]
        title_row = next((row[0].row for row in sheet if row[0].value == group_title), None)
        if title_row is None:
            failures.append(f"Excel lacks {group_title} source values")
        else:
            excel_totals = {}
            for row in sheet.iter_rows(min_row=title_row + 2, min_col=1, max_col=2):
                label, value = row[0].value, row[1].value
                if not isinstance(value, int | float) or not isinstance(label, str):
                    break
                excel_totals[label] = Decimal(str(value))
            if excel_totals.keys() != displayed_totals.keys() or any(
                abs(excel_totals.get(key, Decimal(0)) - expected) > Decimal("0.01")
                for key, expected in displayed_totals.items()
            ):
                failures.append(f"Excel {group_title} values differ from source")
    result = {
        "dataset": name, "source": str(path), "source_facts": source,
        "zephyr_rows": computed["row_count"],
        "zephyr_total": str(actual_total) if actual_total is not None else None,
        "zephyr_blanks": frame[measure_key].null_count() if measure_key in frame.columns else None,
        "zephyr_group_totals": {key: str(value) for key, value in displayed_totals.items()},
        "record_context": context,
        "forecasts": len(computed["forecasts"]),
        "chart_titles": [chart.get("title") for chart in computed["charts"]],
        "pdf_path": str(pdf_path), "pdf_bytes": len(pdf_bytes),
        "pdf_pages": len(page_streams), "export_summary": summary,
        "pdf_chart_titles": [chart.title for chart in visible_charts
                             if chart.title.encode("latin-1", "replace") in pdf_content],
        "xlsx_path": str(xlsx_path), "excel_chart_count": exported_chart_count,
        "kpi_columns": [kpi.get("column") for kpi in computed["kpis"]],
        "executive_summary": computed["ai_result"].get("executive_summary") if with_ai else None,
        "insights": computed["ai_result"].get("insights") if with_ai else None,
        "recommendations": computed["recommendations"],
        "ai_generation": computed["ai_result"].get("generation") if with_ai else None,
        "passed": not failures, "failures": failures,
    }
    return result


async def main(coffee: Path, supermarket: Path, *, with_ai: bool) -> int:
    print("Preparing Kaggle evaluation", file=sys.stderr, flush=True)
    results = []
    workdir = Path(__file__).resolve().parents[3] / "tmp" / "kaggle" / "eval_artifacts"
    workdir.mkdir(parents=True, exist_ok=True)
    for name, path in (("coffee", coffee), ("supermarket", supermarket)):
        try:
            results.append(await run_case(name, path.resolve(), workdir, with_ai=with_ai))
        except Exception as exc:
            results.append({"dataset": name, "passed": False,
                            "failures": [f"{type(exc).__name__}: {exc}"]})
    print(json.dumps({"offline": not with_ai, "results": results}, indent=2, default=str))
    return int(any(not result["passed"] for result in results))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--coffee", type=Path, required=True)
    parser.add_argument("--supermarket", type=Path, required=True)
    parser.add_argument("--with-ai", action="store_true", help="Include live AI wording checks (uses API credits)")
    args = parser.parse_args()
    print("Parsed Kaggle evaluation arguments", file=sys.stderr, flush=True)
    os.environ["DEBUG"] = "false"
    _bootstrap_env(not args.with_ai)
    logging.getLogger("httpx").setLevel(logging.WARNING)
    print("Bootstrapped Kaggle evaluation", file=sys.stderr, flush=True)
    raise SystemExit(asyncio.run(main(args.coffee, args.supermarket, with_ai=args.with_ai)))
