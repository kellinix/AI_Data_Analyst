"""Exported charts retain plotted values in PDF and editable Excel formats."""

from __future__ import annotations

import base64
import io
import re
import zipfile
import zlib

from openpyxl import load_workbook

from app.services.export_charts import normalise_charts
from app.services.export_service import build_excel_report, build_pdf_report


def _charts() -> list[dict]:
    return [
        {"type": "line", "title": "Revenue over time", "xAxis": "Week", "yAxis": "Revenue",
         "echarts_option": {"xAxis": {"data": ["Jan", "Feb", "Mar"]},
                             "series": [{"name": "Revenue", "data": [10, 20, 15]}]}},
        {"type": "bar", "title": "Revenue by region", "xAxis": "Revenue", "yAxis": "Region",
         "echarts_option": {"yAxis": {"type": "category", "data": ["North", "South"]},
                             "series": [{"name": "Revenue", "data": [30, 40]}]}},
        {"type": "bar", "title": "Orders by region and status", "series_by": "Status",
         "xAxis": "Region", "yAxis": "Orders", "aggregation": "sum",
         "echarts_option": {"xAxis": {"data": ["North", "South"]},
                             "series": [{"name": "Open", "data": [3, 5]},
                                        {"name": "Closed", "data": [2, 4]}]}},
        {"type": "donut", "title": "Status breakdown",
         "echarts_option": {"series": [{"data": [
             {"name": "Open", "value": 7}, {"name": "Closed", "value": 9}]}]}},
        {"type": "scatter", "title": "Spend vs revenue", "xAxis": "Spend", "yAxis": "Revenue",
         "echarts_option": {"series": [{"data": [[1, 2], [3, 6], [4, 9]]}]}},
        {"type": "histogram", "title": "Order size distribution", "xAxis": "Size",
         "echarts_option": {"xAxis": {"data": ["0-10", "10-20"]},
                             "series": [{"name": "Orders", "data": [11, 8]}]}},
        {"type": "line", "title": "Sparse readings", "xAxis": "Week", "yAxis": "Value",
         "echarts_option": {"xAxis": {"data": ["Jan", "Feb", "Mar"]},
                             "series": [{"name": "Value", "data": [3, None, 7],
                                         "lineStyle": {"opacity": 0}}]}},
    ]


def test_all_saved_chart_types_export_as_editable_excel_charts():
    workbook_bytes = build_excel_report(
        name="Sample", row_count=42, column_count=4, summary="Sample summary",
        insights=[], charts=_charts(),
    )
    workbook = load_workbook(io.BytesIO(workbook_bytes))
    assert {"Overview", "Charts", "Chart Data"} <= set(workbook.sheetnames)
    assert len(workbook["Charts"]._charts) == 6
    values = [cell.value for row in workbook["Chart Data"] for cell in row]
    assert {"Revenue over time", "North", "Closed", "Mar", 40, 9} <= set(values)
    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as archive:
        chart_names = [name for name in archive.namelist() if name.startswith("xl/charts/chart") and name.endswith(".xml")]
        assert len(chart_names) == 6
        assert any(b"Sparse readings" in archive.read(name) and b"<a:noFill/>" in archive.read(name)
                   for name in chart_names)


def test_pdf_chart_section_adds_vector_visuals_without_dropping_narrative():
    kpis = [{"type": "summary", "title": f"Metric {index}", "data": {"value": index}}
            for index in range(6)]
    plain = build_pdf_report(name="Sample", row_count=42, column_count=4,
                             summary="Sample summary", insights=kpis)
    illustrated = build_pdf_report(name="Sample", row_count=42, column_count=4,
                                   summary="Sample summary", insights=kpis, charts=_charts())
    assert illustrated.startswith(b"%PDF-")
    assert len(illustrated) > len(plain) + 3_000
    assert len(re.findall(rb"/Type\s*/Page\b", illustrated)) >= 2
    page_contents = [
        zlib.decompress(base64.a85decode(stream.strip(), adobe=True))
        for stream in re.findall(rb"stream\n(.*?)endstream", illustrated, re.S)
    ]
    for chart in _charts():
        if chart["type"] == "histogram":
            assert not any(chart["title"].encode() in page for page in page_contents)
            continue
        assert any(chart["title"].encode() in page for page in page_contents)
    assert b"Visual Charts" in page_contents[0]
    assert b"Revenue over time" in page_contents[0]


def test_invalid_or_unpopulated_charts_are_excluded_without_invented_values():
    charts = _charts() + [
        {"type": "line", "title": "Empty", "echarts_option": {"xAxis": {"data": ["Jan"]},
                                                      "series": [{"data": [None]}]}},
        {"type": "scatter", "title": "Bad", "echarts_option": {"series": [{"data": [[1, "nan"]]}]}},
    ]
    assert len(normalise_charts(charts)) == 7


def test_exports_do_not_invent_us_dollars_for_unknown_currency():
    insights = [{"type": "summary", "title": "Sales Amount",
                 "data": {"value": 115431.58, "is_currency": True, "currency": None}}]
    args = {"name": "Coffee", "row_count": 3636, "column_count": 6,
            "summary": None, "insights": insights}
    pdf = build_pdf_report(**args)
    page_content = b"\n".join(
        zlib.decompress(base64.a85decode(stream.strip(), adobe=True))
        for stream in re.findall(rb"stream\n(.*?)endstream", pdf, re.S)
    )
    assert b"115,431.58" in page_content
    assert b"$115,431.58" not in page_content
    workbook = load_workbook(io.BytesIO(build_excel_report(**args)))
    assert workbook["Key Metrics"]["B2"].value == 115431.58
    assert "$" not in workbook["Key Metrics"]["B2"].number_format
